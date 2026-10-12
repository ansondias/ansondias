"""Generate the ``Student Knowledge Base`` documents from Anson Dias's real CV.

Run once to (re)create the personal documents the RAG system indexes:

    python build_knowledge_base.py

Formats produced: Markdown (.md), plain text (.txt), PDF (.pdf), Word (.docx).

The content is taken from Anson Dias's actual resume/portfolio. The academic
record and internship documents are the student's own honest records (no forged
official seals or signatures) presented for the knowledge base.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors

KB = Path(__file__).resolve().parent / "Student Knowledge Base"
SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


# ---------------------------------------------------------------------------
# 1. Resume (Markdown)
# ---------------------------------------------------------------------------
RESUME_MD = """# Anson Dias — Resume

## Contact
- **Full name:** Anson Dias
- **Email:** ansondias18@gmail.com
- **Phone:** +91 90829 24744
- **Location:** Mumbai, India
- **GitHub:** https://github.com/ansondias
- **LinkedIn:** https://www.linkedin.com/in/anson-dias-6411072a1/

## Profile
4th-year B.Tech (Computer Science, AI & ML) student at Atlas SkillTech
University with hands-on experience across ML pipelines, NLP, blockchain, and
operational analytics. Combines technical depth with real business context
through ongoing operations and accounting work across three family-run
enterprises in manufacturing and engineering.

## Education
- **Atlas SkillTech University, Mumbai** — B.Tech in Computer Science
  (AI & Machine Learning), 2023–2027. **GPA: 6.70 / 10** (through 3rd year).
  - Relevant coursework: Machine Learning, Neural Networks, Blockchain
    Development, Cybersecurity, Data Structures & Algorithms, Human-Computer
    Interaction, Distributed Systems.
- **Hiranandani Foundation School, Powai, Mumbai** — ISC (Class XI & XII),
  Science Stream, 2021–2023. **76% aggregate.**

## Technical skills
- **Programming languages:** Python, JavaScript, Solidity, SQL, C++
- **ML / AI:** Scikit-learn, Pandas, NumPy, NLP, Jupyter
- **Web & Blockchain:** React, HTML/CSS, Web3.js, IPFS, Netlify
- **Data & BI:** Power BI, Excel, DAX
- **Tools:** Figma, Git, VS Code
- **Business / Ops:** Zoho Books, Tally, Invoice Management, Bookkeeping

## Experience
- **AI Content & Prompt Engineering Intern — CostumePeti (Startup), esaapa.com (2024)**
  - Engineered custom prompts for the Fashion AI Virtual Try-On feature,
    improving AI image generation quality and output consistency.
  - Generated AI-based fashion imagery and assisted the product team with
    backend prompt configuration for the virtual try-on pipeline.
  - Iterated rapidly on prompt outputs based on user feedback.
- **Operations & Analytics — Pinnacle Enterprises (Powder Coating, SIPCOT Hosur), Apr 2025 – Present**
  - Built an integrated cost-tracking workbook modelling per-sq-m unit economics
    across 7-stage pretreatment, 3-booth manual spray coating and gas-curing.
  - Designed a production tracking system covering ~9,100 sq m/month throughput
    and ~10 oven cycles/day, benchmarking 95% powder transfer efficiency.
  - Manage the accounting workflow via Zoho Books (invoicing, AR ageing, monthly close).
- **Accounts & Bookkeeping — Anson Engg Works & Icarus Enterprises (Automotive Engineering Manufacturing), Ongoing**
  - End-to-end accounting workflow: invoice entry, bills processing, vendor
    reconciliation and bookkeeping; maintaining financial records.

## Projects
- **Cost Analytics Workbook — Powder Coating SME** (Pinnacle Enterprises, 2025):
  integrated multi-sheet workbook modelling per-sq-m unit economics. Stack: Excel,
  structured workbook design, unit-economics modelling.
- **Multi-Business Accounting Automation:** Python file-handling scripts that
  automate extraction and processing of invoice and bill documents across three
  family-business accounting workflows. Stack: Python, file handling, document
  parsing, openpyxl.
- **Blockchain DApp — Decentralised User Registry:** Ethereum DApp storing user
  data via Solidity smart contracts with a Web3.js frontend and IPFS storage.
  Stack: Solidity, Web3.js, Ethereum, IPFS.
- **Pinnacle Enterprises — Website Development & Deployment:** client-facing
  website for a powder-coating SME (Honda Parts, TVS Motors, Sundaram Auto),
  Formspree contact handling, deployed via Netlify with CI on push.
  Stack: HTML/CSS, JavaScript, Netlify, Formspree.
- **ML Specialization Capstone** *(final-year / capstone project, in progress, Q2 2026)*:
  DeepLearning.AI Machine Learning Specialization capstone projects (supervised
  learning, advanced algorithms, unsupervised learning + recommenders),
  implemented in **Python** (Jupyter, scikit-learn).

## Final-year / capstone project language
The final-year / capstone project (ML Specialization Capstone) is implemented in **Python** (Jupyter, scikit-learn).

## Certifications
- Google Machine Learning Crash Course — in progress (Q2 2026).
- DeepLearning.AI Machine Learning Specialization (Coursera) — in progress (Q2 2026).

## Languages
English (Fluent), Hindi (Fluent), Marathi (Basic), French (Basic).
"""


# ---------------------------------------------------------------------------
# 2. Certifications (plain text)
# ---------------------------------------------------------------------------
CERTS_TXT = """CERTIFICATIONS AND COURSES — ANSON DIAS
========================================

1. Google Machine Learning Crash Course — in progress (expected Q2 2026).
2. DeepLearning.AI Machine Learning Specialization (Coursera) — in progress
   (expected Q2 2026). Covers supervised learning, advanced algorithms, and
   unsupervised learning with recommender systems.

Focus areas: machine learning, supervised and unsupervised learning,
recommender systems, and applied ML with Python.
"""


# ---------------------------------------------------------------------------
# 3. Achievements and highlights (Markdown)
# ---------------------------------------------------------------------------
ACHIEVEMENTS_MD = """# Key Achievements & Highlights — Anson Dias

- Built and **deployed a client-facing company website** for a powder-coating SME
  (Honda Parts, TVS Motors, Sundaram Auto) with Netlify CI-on-push and Formspree
  contact handling.
- Designed a production tracking system benchmarking **95% powder transfer
  efficiency** across ~9,100 sq m/month throughput and ~10 oven cycles/day.
- Built an integrated **cost-tracking workbook** modelling per-sq-m unit economics
  across a 7-stage pretreatment and 3-booth coating line, enabling quoting
  accuracy and per-order margin visibility.
- Automated **multi-business accounting** with Python document-parsing scripts,
  cutting manual data entry during periodic close cycles across three enterprises.
- Developed a **blockchain DApp** (Ethereum / Solidity / Web3.js / IPFS) for a
  decentralised user registry.
- Selected as an **AI Content & Prompt Engineering Intern** at CostumePeti,
  improving a Fashion AI Virtual Try-On image pipeline.
"""


# ---------------------------------------------------------------------------
# 4. Project report (Markdown)
# ---------------------------------------------------------------------------
PROJECT_REPORT_MD = """# Project Report — Multi-Business Accounting Automation

**Author:** Anson Dias

## Objective
Reduce manual data-entry effort during periodic accounting close cycles by
automatically extracting and processing invoice and bill documents across three
family-run businesses.

## Method
- Wrote **Python** file-handling scripts to read and parse invoice/bill documents.
- Extracted key fields and structured them for the accounting workflow.
- Integrated the output with the Zoho Books workflow used at Pinnacle Enterprises.

## Tools / Stack
Python, file handling, document parsing, openpyxl.

## Outcome
- Cut manual data-entry effort during monthly close across three enterprises.
- Made invoice/bill processing more consistent and repeatable.

## Related work
A separate cost-analytics workbook models per-sq-m unit economics for the powder
coating business (Excel, unit-economics modelling), and a client-facing website
was deployed on Netlify with CI on push.
"""


def write_text_files() -> None:
    (KB / "resume.md").write_text(RESUME_MD, encoding="utf-8")
    (KB / "certifications.txt").write_text(CERTS_TXT, encoding="utf-8")
    (KB / "achievements_and_awards.md").write_text(ACHIEVEMENTS_MD, encoding="utf-8")
    (KB / "project_report.md").write_text(PROJECT_REPORT_MD, encoding="utf-8")


# ---------------------------------------------------------------------------
# 5. Academic record / transcript (PDF)
# ---------------------------------------------------------------------------
def write_transcript_pdf() -> None:
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=16)

    doc = SimpleDocTemplate(str(KB / "academic_transcript.pdf"), pagesize=A4,
                            topMargin=20 * mm, bottomMargin=20 * mm)
    story = [
        Paragraph("Atlas SkillTech University, Mumbai", title),
        Paragraph("Academic Record / Transcript (through 3rd year)", styles["Heading2"]),
        Spacer(1, 6 * mm),
        Paragraph("<b>Student name:</b> Anson Dias", styles["Normal"]),
        Paragraph("<b>Programme:</b> B.Tech — Computer Science (AI &amp; Machine Learning)", styles["Normal"]),
        Paragraph("<b>Duration:</b> 2023 – 2027 (currently 4th year)", styles["Normal"]),
        Paragraph("<b>GPA:</b> 6.70 / 10 (through 3rd year)", styles["Normal"]),
        Spacer(1, 5 * mm),
        Paragraph("<b>Relevant coursework</b>", styles["Heading3"]),
    ]
    rows = [
        ["Area", "Courses"],
        ["Core CS", "Data Structures & Algorithms, Distributed Systems"],
        ["AI / ML", "Machine Learning, Neural Networks"],
        ["Applied", "Blockchain Development, Cybersecurity, Human-Computer Interaction"],
    ]
    table = Table(rows, colWidths=[30 * mm, 120 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2d3748")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story += [
        table,
        Spacer(1, 6 * mm),
        Paragraph("<b>Final-year / capstone project:</b> ML Specialization Capstone "
                  "(in progress, Q2 2026), implemented in Python (Jupyter, scikit-learn).",
                  styles["Normal"]),
        Spacer(1, 5 * mm),
        Paragraph("<b>Prior education:</b> ISC (Class XI &amp; XII), Science Stream, "
                  "Hiranandani Foundation School, Powai, Mumbai, 2021–2023 — 76% aggregate.",
                  styles["Normal"]),
    ]
    doc.build(story)


# ---------------------------------------------------------------------------
# 6. Internship experience certificate (DOCX)
# ---------------------------------------------------------------------------
def write_internship_docx() -> None:
    d = DocxDocument()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    d.add_heading("Internship Experience", level=0)
    d.add_heading("AI Content & Prompt Engineering Intern — CostumePeti", level=2)
    d.add_paragraph("Organisation: CostumePeti (startup) · esaapa.com")
    d.add_paragraph("Role: AI Content & Prompt Engineering Intern · Year: 2024")
    d.add_paragraph(
        "Anson Dias worked as an AI Content & Prompt Engineering Intern at "
        "CostumePeti in 2024. During the internship he engineered custom prompts "
        "for the Fashion AI Virtual Try-On feature, improving AI image generation "
        "quality and output consistency."
    )
    d.add_paragraph(
        "He generated AI-based fashion imagery, assisted the product team with "
        "backend prompt configuration for the virtual try-on pipeline, and iterated "
        "rapidly on prompt outputs based on user feedback in a fast-paced startup "
        "environment."
    )
    d.save(str(KB / "internship_certificate.docx"))


# ---------------------------------------------------------------------------
# 7. Statement of Purpose (DOCX)
# ---------------------------------------------------------------------------
def write_sop_docx() -> None:
    d = DocxDocument()
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading("Statement of Purpose", level=0)
    d.add_heading("Anson Dias", level=2)
    d.add_paragraph(
        "My name is Anson Dias and I am a 4th-year B.Tech student in Computer "
        "Science (AI & Machine Learning) at Atlas SkillTech University, Mumbai."
    )
    d.add_paragraph(
        "My interests lie in machine learning pipelines, natural language "
        "processing, blockchain development, and operational / business analytics. "
        "I enjoy applying AI and data to real-world problems — I run operations, "
        "analytics and accounting workflows across three family-run manufacturing "
        "and engineering businesses, and I build the tooling that supports them."
    )
    d.add_paragraph(
        "I have engineered prompts for a Fashion AI virtual try-on pipeline, "
        "automated multi-business accounting with Python, built a blockchain DApp, "
        "and deployed a client-facing website. Going forward I want to deepen my "
        "work in machine learning and retrieval-based systems that ground AI in "
        "trustworthy, domain-specific data."
    )
    d.add_paragraph(
        "I can be reached at ansondias18@gmail.com or +91 90829 24744, and via "
        "GitHub (github.com/ansondias) and LinkedIn."
    )
    d.save(str(KB / "statement_of_purpose.docx"))


def main() -> None:
    KB.mkdir(parents=True, exist_ok=True)
    # Remove any previous knowledge-base files so no stale documents remain.
    for p in KB.iterdir():
        if p.is_file() and p.suffix.lower() in SUPPORTED:
            p.unlink()

    write_text_files()
    write_transcript_pdf()
    write_internship_docx()
    write_sop_docx()

    files = sorted(p.name for p in KB.iterdir() if p.is_file())
    print(f"Wrote {len(files)} documents to '{KB.name}/':")
    for name in files:
        size = (KB / name).stat().st_size
        print(f"  - {name:32s} ({size:,} bytes)")


if __name__ == "__main__":
    main()
