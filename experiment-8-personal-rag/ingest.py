"""Build the vector knowledge base.

Pipeline:  load documents  ->  split into chunks  ->  embed  ->  store in Chroma.

Run:
    python ingest.py
"""
from __future__ import annotations

import shutil
from pathlib import Path
from typing import List

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma

import config
from embeddings import LocalONNXEmbeddings


# ---------------------------------------------------------------------------
# Document loading (PDF / DOCX / TXT / MD)
# ---------------------------------------------------------------------------
def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


_LOADERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".txt": _read_text,
    ".md": _read_text,
}


def load_documents() -> List[Document]:
    """Read every supported file in the knowledge base into LangChain Documents."""
    docs: List[Document] = []
    paths = sorted(config.KNOWLEDGE_BASE_DIR.glob("*"))
    for path in paths:
        if not path.is_file() or path.suffix.lower() not in config.SUPPORTED_EXTENSIONS:
            continue
        loader = _LOADERS[path.suffix.lower()]
        text = loader(path).strip()
        if not text:
            continue
        docs.append(
            Document(
                page_content=text,
                metadata={"source": path.name, "type": path.suffix.lower().lstrip(".")},
            )
        )
        print(f"  loaded  {path.name:32s} ({len(text):,} chars)")
    return docs


# ---------------------------------------------------------------------------
# Split -> embed -> store
# ---------------------------------------------------------------------------
def build(reset: bool = True) -> Chroma:
    if reset and config.CHROMA_DIR.exists():
        shutil.rmtree(config.CHROMA_DIR)

    print(f"[1/4] Loading documents from '{config.KNOWLEDGE_BASE_DIR.name}/'")
    documents = load_documents()
    if not documents:
        raise SystemExit("No documents found. Run build_knowledge_base.py first.")

    print(f"[2/4] Splitting {len(documents)} documents into chunks "
          f"(size={config.CHUNK_SIZE}, overlap={config.CHUNK_OVERLAP})")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE,
        chunk_overlap=config.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(documents)
    print(f"      -> {len(chunks)} chunks")

    print("[3/4] Embedding chunks with all-MiniLM-L6-v2 (384-dim ONNX, CPU)")
    embeddings = LocalONNXEmbeddings()

    print(f"[4/4] Storing vectors in Chroma at '{config.CHROMA_DIR.name}/'")
    store = Chroma.from_documents(
        documents=chunks,
        embedding=embeddings,
        collection_name=config.COLLECTION_NAME,
        persist_directory=str(config.CHROMA_DIR),
        collection_metadata={"hnsw:space": "cosine"},
    )
    print(f"\nDone. Indexed {len(chunks)} chunks from {len(documents)} documents.")
    return store


if __name__ == "__main__":
    build()
